import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference


file1 = None
file2 = None


def choose_file1():
    global file1
    file1 = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
    if file1:
        btn1.config(text="1-й файл выбран")


def choose_file2():
    global file2
    file2 = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
    if file2:
        btn2.config(text="2-й файл выбран")


def generate_excel():
    if not file1 or not file2:
        messagebox.showerror("Ошибка", "Выберите оба файла")
        return

    # читаем CSV
    df1 = pd.read_csv(file1)
    df2 = pd.read_csv(file2)

    # ожидаем столбцы x,y
    x = df1.iloc[:, 0]
    y1 = df1.iloc[:, 1]
    y2 = df2.iloc[:, 1]

    diff = y1 - y2

    # создаём Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    ws.append(["x", "y1", "y2", "y1 - y2"])

    for i in range(len(x)):
        ws.append([x[i], y1[i], y2[i], diff[i]])

    # -----------------------------
    # Диаграмма 1: y1 и y2 от x
    # -----------------------------
    chart1 = LineChart()
    chart1.title = "y1 и y2 от x"
    chart1.x_axis.title = "x"
    chart1.y_axis.title = "y"

    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(x) + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(x) + 1)

    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)

    ws.add_chart(chart1, "F2")

    # -----------------------------
    # Диаграмма 2: (y1 - y2) от x
    # -----------------------------
    chart2 = LineChart()
    chart2.title = "y1 - y2 от x"
    chart2.x_axis.title = "x"
    chart2.y_axis.title = "Δy"

    data2 = Reference(ws, min_col=4, min_row=1, max_row=len(x) + 1)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats)

    ws.add_chart(chart2, "F20")

    # сохраняем
    wb.save("result.xlsx")
    messagebox.showinfo("Готово", "Файл result.xlsx создан")


# -----------------------------
# GUI
# -----------------------------
root = tk.Tk()
root.title("Генератор Excel диаграмм")
root.geometry("300x200")

btn1 = tk.Button(root, text="Выбрать 1-й файл", command=choose_file1)
btn1.pack(pady=10)

btn2 = tk.Button(root, text="Выбрать 2-й файл", command=choose_file2)
btn2.pack(pady=10)

btn3 = tk.Button(root, text="Сгенерировать", command=generate_excel)
btn3.pack(pady=20)

root.mainloop()


# # Press the green button in the gutter to run the script.
# if __name__ == '__main__':
#     print_hi('PyCharm')

# See PyCharm help at https://www.jetbrains.com/help/pycharm/
